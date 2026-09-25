import openpyxl, csv, os

for excelFile in os.listdir('.'):
    # Skip non-xlsx files
    if not excelFile.endswith('.xlsx'):
        continue

    wb = openpyxl.load_workbook(excelFile)
    for sheetName in wb.sheetnames:
        sheet = wb[sheetName]

        # Create CSV filename
        excelBase = os.path.splitext(excelFile)[0]
        csvFilename = excelBase + '_' + sheetName + '.csv'

        # Create CSV writer
        csvFile = open(csvFilename, 'w', newline='')
        csvWriter = csv.writer(csvFile)

        # Loop through every row
        for rowNum in range(1, sheet.max_row + 1):
            rowData = []
            # Loop through every column in this row
            for colNum in range(1, sheet.max_column + 1):
                cell = sheet.cell(row=rowNum, column=colNum)
                rowData.append(cell.value)
            csvWriter.writerow(rowData)

        csvFile.close()