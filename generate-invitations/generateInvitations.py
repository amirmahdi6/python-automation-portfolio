import docx

doc = docx.Document()
with open('guests.txt') as f:
    names = f.readlines()

for name in names:
    name = name.strip()
    doc.add_paragraph('It would be a pleasure to have the company of', 'Normal')
    doc.add_paragraph(name, 'Normal')
    doc.add_paragraph('at 11010 Memory Lane on the Evening of', 'Normal')
    doc.add_paragraph('April 1st', 'Normal')
    doc.add_paragraph("at 7 o'clock", 'Normal')
    doc.paragraphs[-1].runs[0].add_break(docx.enum.text.WD_BREAK.PAGE)

doc.save('invitations.docx')